from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask import Blueprint, g, jsonify, request, send_file
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError

if __package__:
    from .auth import require_auth, require_roles
    from .database import db
    from .models import (
        Attendance,
        Course,
        Exam,
        ExamResult,
        FeePayment,
        FeeRecord,
        Student,
    )
else:
    from auth import require_auth, require_roles
    from database import db
    from models import Attendance, Course, Exam, ExamResult, FeePayment, FeeRecord, Student

modules_api = Blueprint("modules_api", __name__, url_prefix="/api")
ATTENDANCE_STATUSES = {"Present", "Absent", "Late", "Excused"}


def parse_date(value, field):
    if not isinstance(value, str):
        return None, "{} must use YYYY-MM-DD format.".format(field)
    try:
        return date.fromisoformat(value), None
    except ValueError:
        return None, "{} must use YYYY-MM-DD format.".format(field)


def student_for_user(student_id):
    if g.current_user.role == "student":
        if not g.current_user.student_id:
            return None, (jsonify({"error": "This student account is not linked to a student record."}), 403)
        if student_id not in (None, "", g.current_user.student_id, str(g.current_user.student_id)):
            return None, (jsonify({"error": "You do not have permission to access that record."}), 403)
        student_id = g.current_user.student_id
    if student_id in (None, ""):
        return None, None
    try:
        student_id = int(student_id)
    except (TypeError, ValueError):
        return None, (jsonify({"error": "Student ID must be a number."}), 400)
    student = db.session.get(Student, student_id)
    if not student:
        return None, (jsonify({"error": "Student not found."}), 404)
    return student, None


def grade_for(marks, total):
    percentage = marks * 100.0 / total
    if percentage >= 90:
        return "A"
    if percentage >= 80:
        return "B"
    if percentage >= 70:
        return "C"
    if percentage >= 60:
        return "D"
    return "F"


def fee_dict(fee):
    paid = db.session.query(func.coalesce(func.sum(FeePayment.amount), 0)).filter_by(
        fee_id=fee.id
    ).scalar()
    total = Decimal(fee.amount)
    paid = Decimal(paid)
    return {
        "id": fee.id,
        "student_id": fee.student_id,
        "student_name": db.session.get(Student, fee.student_id).full_name,
        "description": fee.description,
        "amount": float(total),
        "paid": float(paid),
        "pending": float(total - paid),
        "status": "Paid" if paid >= total else ("Partial" if paid else "Pending"),
        "due_date": fee.due_date.isoformat(),
        "payments": [
            {
                "id": payment.id,
                "amount": float(payment.amount),
                "paid_at": payment.paid_at.isoformat(),
                "reference": payment.reference,
            }
            for payment in FeePayment.query.filter_by(fee_id=fee.id).order_by(FeePayment.paid_at).all()
        ],
    }


@modules_api.get("/courses")
@require_auth
def list_courses():
    registered = [course.to_dict() for course in Course.query.order_by(Course.name).all()]
    names = {course["name"].lower() for course in registered}
    for name, department in db.session.query(Student.course, Student.department).distinct().all():
        if name.lower() not in names:
            registered.append({"id": None, "code": "", "name": name, "department": department, "credits": None})
    return jsonify({"courses": registered})


@modules_api.post("/courses")
@require_roles("admin")
def create_course():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON object is required."}), 400
    code = data.get("code")
    name = data.get("name")
    department = data.get("department")
    credits = data.get("credits", 3)
    if not all(isinstance(value, str) and value.strip() for value in (code, name, department)):
        return jsonify({"error": "Course code, name, and department are required."}), 400
    if any(len(value.strip()) > limit for value, limit in ((code, 30), (name, 120), (department, 100))):
        return jsonify({"error": "Course fields exceed their maximum length."}), 400
    if not isinstance(credits, int) or isinstance(credits, bool) or credits < 1 or credits > 30:
        return jsonify({"error": "Credits must be a number from 1 to 30."}), 400
    course = Course(code=code.strip().upper(), name=name.strip(), department=department.strip(), credits=credits)
    db.session.add(course)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "A course with this code already exists."}), 409
    return jsonify({"message": "Course created.", "course": course.to_dict()}), 201


@modules_api.get("/attendance")
@require_auth
def attendance_history():
    query = Attendance.query
    student_id = request.args.get("student_id")
    student, error = student_for_user(student_id)
    if error:
        return error
    if student:
        query = query.filter_by(student_id=student.id)
    selected_date = request.args.get("date")
    if selected_date:
        parsed, error_message = parse_date(selected_date, "Date")
        if error_message:
            return jsonify({"error": error_message}), 400
        query = query.filter_by(attendance_date=parsed)
    records = query.order_by(Attendance.attendance_date.desc(), Attendance.id.desc()).limit(1000).all()
    result = []
    for record in records:
        row = record.to_dict()
        row["student_name"] = db.session.get(Student, record.student_id).full_name
        result.append(row)
    return jsonify({"attendance": result})


@modules_api.get("/student/dashboard")
@require_roles("student")
def student_dashboard():
    student, error = student_for_user(None)
    if error:
        return error
    attendance_query = Attendance.query.filter_by(student_id=student.id)
    records = attendance_query.order_by(
        Attendance.attendance_date.desc(), Attendance.id.desc()
    ).limit(1000).all()
    total = attendance_query.count()
    attended = attendance_query.filter(Attendance.status.in_(("Present", "Late"))).count()
    absent = attendance_query.filter_by(status="Absent").count()
    return jsonify({
        "student": student.to_dict(),
        "attendance": [record.to_dict() for record in records],
        "summary": {
            "total_classes": total,
            "classes_attended": attended,
            "classes_absent": absent,
            "attendance_percent": round(attended * 100.0 / total, 1) if total else 0,
        },
    })


@modules_api.get("/attendance/students/<int:student_id>")
@require_auth
def student_attendance_history(student_id):
    student, error = student_for_user(str(student_id))
    if error:
        return error
    records = Attendance.query.filter_by(student_id=student.id).order_by(
        Attendance.attendance_date.desc(), Attendance.id.desc()
    ).limit(1000).all()
    return jsonify({
        "student": student.to_dict(),
        "attendance": [record.to_dict() for record in records],
    })


@modules_api.get("/attendance/report")
@require_roles("admin", "faculty")
def attendance_report():
    query = Attendance.query
    start_value = request.args.get("start_date")
    end_value = request.args.get("end_date")
    start_date = end_date = None
    if start_value:
        start_date, error = parse_date(start_value, "Start date")
        if error:
            return jsonify({"error": error}), 400
        query = query.filter(Attendance.attendance_date >= start_date)
    if end_value:
        end_date, error = parse_date(end_value, "End date")
        if error:
            return jsonify({"error": error}), 400
        query = query.filter(Attendance.attendance_date <= end_date)
    if start_date and end_date and start_date > end_date:
        return jsonify({"error": "Start date must be on or before end date."}), 400

    records = query.order_by(Attendance.attendance_date, Attendance.student_id).all()
    daily = {}
    by_student = {}
    for record in records:
        day = record.attendance_date.isoformat()
        day_counts = daily.setdefault(day, {
            "date": day, "total": 0, "present": 0, "absent": 0, "late": 0, "excused": 0,
        })
        student_counts = by_student.setdefault(record.student_id, {
            "student_id": record.student_id,
            "total_classes": 0,
            "classes_attended": 0,
            "classes_absent": 0,
        })
        status = record.status.lower()
        day_counts["total"] += 1
        day_counts[status] += 1
        student_counts["total_classes"] += 1
        if record.status in ("Present", "Late"):
            student_counts["classes_attended"] += 1
        elif record.status == "Absent":
            student_counts["classes_absent"] += 1

    daily_rows = sorted(daily.values(), key=lambda row: row["date"], reverse=True)
    student_rows = []
    for student_id, counts in by_student.items():
        student = db.session.get(Student, student_id)
        if not student:
            continue
        student_rows.append({
            **counts,
            "student": student.to_dict(),
            "attendance_percent": round(
                counts["classes_attended"] * 100.0 / counts["total_classes"], 1
            ),
        })
    student_rows.sort(key=lambda row: row["student"]["full_name"].lower())
    return jsonify({
        "daily": daily_rows,
        "students": student_rows,
        "total_classes": len(daily_rows),
        "total_records": len(records),
    })


@modules_api.post("/attendance")
@require_roles("admin", "faculty")
def mark_attendance():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON object is required."}), 400
    attendance_date, error = parse_date(data.get("date"), "Date")
    if error:
        return jsonify({"error": error}), 400
    if attendance_date > date.today():
        return jsonify({"error": "Attendance cannot be marked for a future date."}), 400
    entries = data.get("records")
    if not isinstance(entries, list) or not entries:
        return jsonify({"error": "Provide at least one attendance record."}), 400

    updated = []
    for entry in entries:
        if not isinstance(entry, dict):
            return jsonify({"error": "Each attendance record must be an object."}), 400
        student_id = entry.get("student_id")
        status = entry.get("status")
        if (not isinstance(student_id, int) or isinstance(student_id, bool)
                or not isinstance(status, str) or status not in ATTENDANCE_STATUSES):
            return jsonify({"error": "Each record needs a numeric student_id and valid status."}), 400
        if not db.session.get(Student, student_id):
            return jsonify({"error": "Student {} was not found.".format(student_id)}), 404
        record = Attendance.query.filter_by(
            student_id=student_id, attendance_date=attendance_date
        ).first()
        if record:
            record.status = status
            record.marked_by = g.current_user.id
        else:
            record = Attendance(
                student_id=student_id,
                attendance_date=attendance_date,
                status=status,
                marked_by=g.current_user.id,
            )
            db.session.add(record)
        updated.append(record)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({"error": "Attendance changed concurrently. Refresh and retry."}), 409
    return jsonify({"message": "Attendance saved.", "count": len(updated)})


@modules_api.get("/attendance/summary")
@require_auth
def attendance_summary():
    student, error = student_for_user(request.args.get("student_id"))
    if error:
        return error
    query = Attendance.query
    if student:
        query = query.filter_by(student_id=student.id)
    total = query.count()
    present = query.filter(Attendance.status.in_(("Present", "Late"))).count()
    return jsonify({
        "total_days": total,
        "present_days": present,
        "absent_days": query.filter_by(status="Absent").count(),
        "attendance_percent": round(present * 100.0 / total, 1) if total else 0,
    })


def exam_results(exam):
    results = ExamResult.query.filter_by(exam_id=exam.id).order_by(ExamResult.student_id).all()
    return [
        {
            "id": result.id,
            "student_id": result.student_id,
            "student_name": db.session.get(Student, result.student_id).full_name,
            "marks": result.marks,
            "total_marks": exam.total_marks,
            "grade": result.grade,
            "percentage": round(result.marks * 100.0 / exam.total_marks, 1),
        }
        for result in results
    ]


@modules_api.get("/exams")
@require_auth
def list_exams():
    exams = Exam.query.order_by(Exam.exam_date.desc(), Exam.id.desc()).all()
    student, error = student_for_user(request.args.get("student_id"))
    if error:
        return error
    output = []
    for exam in exams:
        results = exam_results(exam)
        if student:
            results = [result for result in results if result["student_id"] == student.id]
        output.append(dict(exam.to_dict(), results=results))
    return jsonify({"exams": output})


@modules_api.post("/exams")
@require_roles("admin", "faculty")
def create_exam():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON object is required."}), 400
    title = data.get("title")
    course_name = data.get("course")
    exam_date, error = parse_date(data.get("exam_date"), "Exam date")
    total_marks = data.get("total_marks")
    if not isinstance(title, str) or not title.strip() or len(title.strip()) > 120:
        return jsonify({"error": "Exam title is required (120 characters maximum)."}), 400
    if not isinstance(course_name, str) or not course_name.strip():
        return jsonify({"error": "Course is required."}), 400
    if error:
        return jsonify({"error": error}), 400
    if isinstance(total_marks, bool):
        return jsonify({"error": "Total marks must be a positive number."}), 400
    try:
        total_marks = float(total_marks)
    except (TypeError, ValueError):
        return jsonify({"error": "Total marks must be a positive number."}), 400
    if total_marks <= 0 or total_marks > 100000:
        return jsonify({"error": "Total marks must be between 0 and 100000."}), 400
    course = Course.query.filter(func.lower(Course.name) == course_name.strip().lower()).first()
    exam = Exam(
        title=title.strip(),
        course_name=course.name if course else course_name.strip(),
        course_id=course.id if course else None,
        exam_date=exam_date,
        total_marks=total_marks,
    )
    db.session.add(exam)
    db.session.commit()
    return jsonify({"message": "Examination created.", "exam": exam.to_dict()}), 201


@modules_api.post("/exams/<int:exam_id>/results")
@require_roles("admin", "faculty")
def save_exam_results(exam_id):
    exam = db.session.get(Exam, exam_id)
    if not exam:
        return jsonify({"error": "Examination not found."}), 404
    data = request.get_json(silent=True)
    records = data.get("results") if isinstance(data, dict) else None
    if not isinstance(records, list) or not records:
        return jsonify({"error": "Provide at least one result."}), 400
    for result_data in records:
        if not isinstance(result_data, dict) or not isinstance(result_data.get("student_id"), int):
            return jsonify({"error": "Each result needs a numeric student_id."}), 400
        if isinstance(result_data.get("marks"), bool):
            return jsonify({"error": "Marks must be a number."}), 400
        try:
            marks = float(result_data.get("marks"))
        except (TypeError, ValueError):
            return jsonify({"error": "Marks must be a number."}), 400
        if not db.session.get(Student, result_data["student_id"]):
            return jsonify({"error": "Student {} was not found.".format(result_data["student_id"])}), 404
        if marks < 0 or marks > exam.total_marks:
            return jsonify({"error": "Marks must be between 0 and the exam total."}), 400
        result = ExamResult.query.filter_by(
            exam_id=exam.id, student_id=result_data["student_id"]
        ).first()
        if not result:
            result = ExamResult(exam_id=exam.id, student_id=result_data["student_id"], marks=marks, grade=grade_for(marks, exam.total_marks))
            db.session.add(result)
        else:
            result.marks = marks
            result.grade = grade_for(marks, exam.total_marks)
    db.session.commit()
    return jsonify({"message": "Results saved.", "results": exam_results(exam)})


@modules_api.get("/fees")
@require_roles("admin", "student")
def list_fees():
    query = FeeRecord.query
    student, error = student_for_user(request.args.get("student_id"))
    if error:
        return error
    if student:
        query = query.filter_by(student_id=student.id)
    return jsonify({"fees": [fee_dict(fee) for fee in query.order_by(FeeRecord.due_date.desc()).all()]})


@modules_api.post("/fees")
@require_roles("admin")
def create_fee():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON object is required."}), 400
    student_id = data.get("student_id")
    description = data.get("description")
    due_date, error = parse_date(data.get("due_date"), "Due date")
    try:
        amount = Decimal(str(data.get("amount")))
    except (InvalidOperation, ValueError):
        return jsonify({"error": "Fee amount must be a positive number."}), 400
    if (not isinstance(student_id, int) or isinstance(student_id, bool)
            or not db.session.get(Student, student_id)):
        return jsonify({"error": "Choose an existing student."}), 404
    if not isinstance(description, str) or not description.strip() or len(description.strip()) > 160:
        return jsonify({"error": "Fee description is required (160 characters maximum)."}), 400
    if error:
        return jsonify({"error": error}), 400
    if not amount.is_finite() or amount <= 0 or amount > Decimal("9999999999.99"):
        return jsonify({"error": "Fee amount must be positive and within the supported range."}), 400
    fee = FeeRecord(student_id=student_id, description=description.strip(), amount=amount, due_date=due_date)
    db.session.add(fee)
    db.session.commit()
    return jsonify({"message": "Fee recorded.", "fee": fee_dict(fee)}), 201


@modules_api.post("/fees/<int:fee_id>/payments")
@require_roles("admin")
def record_payment(fee_id):
    fee = db.session.get(FeeRecord, fee_id)
    if not fee:
        return jsonify({"error": "Fee record not found."}), 404
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "A JSON object is required."}), 400
    try:
        amount = Decimal(str(data.get("amount")))
    except (InvalidOperation, ValueError):
        return jsonify({"error": "Payment amount must be a positive number."}), 400
    paid = Decimal(db.session.query(func.coalesce(func.sum(FeePayment.amount), 0)).filter_by(fee_id=fee.id).scalar())
    if not amount.is_finite() or amount <= 0 or amount > Decimal(fee.amount) - paid:
        return jsonify({"error": "Payment must be positive and no greater than the pending balance."}), 400
    reference = data.get("reference", "")
    if not isinstance(reference, str) or len(reference) > 100:
        return jsonify({"error": "Payment reference must be 100 characters or fewer."}), 400
    payment = FeePayment(fee_id=fee.id, amount=amount, reference=reference.strip())
    db.session.add(payment)
    db.session.commit()
    return jsonify({"message": "Payment recorded.", "fee": fee_dict(fee)}), 201


@modules_api.get("/fees/report")
@require_roles("admin")
def fee_report():
    fees = FeeRecord.query.all()
    rows = [fee_dict(fee) for fee in fees]
    return jsonify({
        "total_due": round(sum(row["amount"] for row in rows), 2),
        "total_collected": round(sum(row["paid"] for row in rows), 2),
        "total_pending": round(sum(row["pending"] for row in rows), 2),
        "fees": rows,
    })


@modules_api.get("/dashboard/analytics")
@require_auth
def dashboard_analytics():
    today = date.today()
    current_month = today.replace(day=1)
    months = []
    for offset in range(5, -1, -1):
        month_number = current_month.month - offset
        year = current_month.year
        while month_number <= 0:
            month_number += 12
            year -= 1
        months.append(date(year, month_number, 1))

    next_month_number = current_month.month + 1
    next_year = current_month.year
    if next_month_number == 13:
        next_month_number = 1
        next_year += 1
    end_date = date(next_year, next_month_number, 1)
    month_rows = []
    student_id = g.current_user.student_id if g.current_user.role == "student" else None
    if g.current_user.role == "student" and not student_id:
        return jsonify({"error": "This student account is not linked to a student record."}), 403
    for index, month_start in enumerate(months):
        if index + 1 < len(months):
            month_end = months[index + 1]
        else:
            month_end = end_date
        attendance_query = Attendance.query.filter(
            Attendance.attendance_date >= month_start,
            Attendance.attendance_date < month_end,
        )
        if student_id:
            attendance_query = attendance_query.filter_by(student_id=student_id)
        total = attendance_query.count()
        present = attendance_query.filter(Attendance.status.in_(("Present", "Late"))).count()
        fee_total = 0
        if g.current_user.role == "admin":
            fee_total = float(
                db.session.query(func.coalesce(func.sum(FeePayment.amount), 0)).filter(
                    FeePayment.paid_at >= datetime.combine(month_start, datetime.min.time()),
                    FeePayment.paid_at < datetime.combine(month_end, datetime.min.time()),
                ).scalar()
            )
        month_rows.append({
            "month": month_start.strftime("%b"),
            "attendance_percent": round(present * 100.0 / total, 1) if total else 0,
            "fees_collected": fee_total,
        })
    return jsonify({"months": month_rows})


def pdf_response(title, rows, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), title=title)
    styles = getSampleStyleSheet()
    safe_title = title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    elements = [Paragraph(safe_title, styles["Title"]), Spacer(1, 16)]
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6658d3")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dce0ea")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f6fb")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    elements.append(table)
    document.build(elements)
    output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=filename)


@modules_api.get("/reports/students.xlsx")
@require_roles("admin", "faculty")
def export_students_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    headers = ["Student ID", "Full name", "Email", "Phone", "Gender", "Date of birth", "Course", "Department", "Address", "Enrollment date", "Status"]
    sheet.append(headers)
    for student in Student.query.order_by(Student.full_name).all():
        sheet.append([
            excel_text(student.student_id), excel_text(student.full_name), excel_text(student.email),
            excel_text(student.phone), excel_text(student.gender), student.date_of_birth.isoformat(),
            excel_text(student.course), excel_text(student.department), excel_text(student.address),
            student.enrollment_date.isoformat(), excel_text(student.status),
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6658D3")
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="college-students.xlsx")


def excel_text(value):
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


@modules_api.get("/reports/students.pdf")
@require_roles("admin", "faculty")
def export_students_pdf():
    rows = [["Student ID", "Name", "Email", "Course", "Department", "Status"]]
    for student in Student.query.order_by(Student.full_name).all():
        rows.append([student.student_id, student.full_name, student.email, student.course, student.department, student.status])
    return pdf_response("College student report", rows, "college-students.pdf")


@modules_api.get("/reports/students/<int:student_id>.pdf")
@require_roles("admin", "faculty")
def export_student_pdf(student_id):
    student = db.session.get(Student, student_id)
    if not student:
        return jsonify({"error": "Student not found."}), 404
    rows = [
        ["Field", "Student information"],
        ["Student ID", student.student_id],
        ["Full name", student.full_name],
        ["Email", student.email],
        ["Phone", student.phone or "—"],
        ["Course / department", "{} / {}".format(student.course, student.department)],
        ["Date of birth", student.date_of_birth.isoformat()],
        ["Enrollment date", student.enrollment_date.isoformat()],
        ["Status", student.status],
    ]
    return pdf_response("Student profile: {}".format(student.full_name), rows, "{}-profile.pdf".format(student.student_id.lower()))


@modules_api.get("/reports/fees.pdf")
@require_roles("admin")
def export_fees_pdf():
    rows = [["Student", "Description", "Amount", "Paid", "Pending", "Due date", "Status"]]
    for fee in FeeRecord.query.order_by(FeeRecord.due_date).all():
        record = fee_dict(fee)
        rows.append([
            record["student_name"], record["description"], amount_text(record["amount"]),
            amount_text(record["paid"]), amount_text(record["pending"]),
            record["due_date"], record["status"],
        ])
    return pdf_response("College fee report", rows, "college-fees.pdf")


def amount_text(value):
    return "{:.2f}".format(float(value))
